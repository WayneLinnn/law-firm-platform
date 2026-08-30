# -*- coding: utf-8 -*-
"""从上传的文件里抽正文，供全文检索。

★ **这个文件最重要的一条：抽不出文字时必须说清楚是哪一种「抽不出」。**

律所的案卷里有大量**扫描件** —— 法院送达的裁定书、对方寄来的证据，
拍照或扫描成 PDF，里面一个字符都没有，只有图。

如果这种文件被静默地记成「正文为空」，后果是：律师搜「优先受偿权」，
系统说没找到，他就以为案卷里真没有这份材料 —— 但它明明躺在那儿。
**搜不到比搜错更危险，因为搜不到不会引起怀疑。**

所以分四种状态，界面上要能看见：

    ok           抽到正文了，全文检索能命中
    empty        文件里没有可提取的文字 —— **多半是扫描件**，需要 OCR 才能检索
    unsupported  这个格式我们不解析（图片、压缩包、音视频）
    failed       解析出错（加密、损坏）

OCR 现在**不做**（要装 tesseract 或调云服务，是另一个量级的事）。
不做可以，但必须让律师知道这份文件搜不到。
"""
from __future__ import annotations

import io
import logging
import re
from typing import Tuple

logger = logging.getLogger(__name__)

#: 超过这个长度就截断。全文检索用不着整本书，
#: 而且 Postgres 的行会被 TOAST 到别处，检索反而变慢。
MAX_TEXT = 200_000

#: 我们不解析的格式。列出来是为了给出**准确**的提示，
#: 而不是笼统一句「解析失败」。
BINARY_EXT = {
    ".jpg": "图片", ".jpeg": "图片", ".png": "图片", ".gif": "图片",
    ".bmp": "图片", ".tif": "图片", ".tiff": "图片", ".heic": "图片",
    ".zip": "压缩包", ".rar": "压缩包", ".7z": "压缩包",
    ".mp3": "音频", ".wav": "音频", ".m4a": "音频",
    ".mp4": "视频", ".mov": "视频", ".avi": "视频",
    ".xls": "旧版 Excel", ".ppt": "旧版 PPT", ".doc": "旧版 Word",
}

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log", ".htm", ".html", ".xml"}


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT]


def _ext(filename: str) -> str:
    i = filename.rfind(".")
    return filename[i:].lower() if i >= 0 else ""


def extract(data: bytes, filename: str, content_type: str = "") -> Tuple[str, str, str]:
    """返回 (正文, 状态, 说明)。**永不抛异常** —— 抽取失败不该挡住上传。

    文件本身已经存进对象存储了，抽正文只是为了检索。
    因为一个损坏的 PDF 就让律师传不上材料，是本末倒置。
    """
    ext = _ext(filename)

    if ext in BINARY_EXT:
        return "", "unsupported", "{0}文件不提取文字，全文检索搜不到它".format(BINARY_EXT[ext])

    try:
        if ext == ".pdf" or content_type == "application/pdf":
            return _pdf(data)
        if ext == ".docx":
            return _docx(data)
        if ext in TEXT_EXT or content_type.startswith("text/"):
            return _plain(data)
        # 认不出来的，先当文本试一把 —— 很多东西其实就是文本
        text, status, note = _plain(data)
        if status == "ok" and len(text) > 20:
            return text, status, note
        return "", "unsupported", "无法识别的格式（{0}），全文检索搜不到它".format(ext or "无扩展名")
    except Exception as exc:                      # noqa: BLE001
        logger.warning("抽取 %s 失败：%s", filename, exc)
        return "", "failed", "解析出错：{0}".format(exc)[:200]


def _pdf(data: bytes) -> Tuple[str, str, str]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:                      # noqa: BLE001
        return "", "failed", "PDF 打不开：{0}".format(exc)[:180]

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")                    # 有些只是空口令
        except Exception:                         # noqa: BLE001
            return "", "failed", "PDF 已加密，无法提取文字"

    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:                         # noqa: BLE001
            continue
    text = _clean("\n".join(parts))

    if not text:
        # ★ 这就是扫描件。必须说人话，不能只报「空」
        return "", "empty", (
            "这份 PDF 里没有文字层，**多半是扫描件**。"
            "文件本身已存好，但全文检索搜不到它 —— 需要 OCR。"
        )
    return text, "ok", ""


def _docx(data: bytes) -> Tuple[str, str, str]:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # 表格里常放证据清单、结算明细，不能漏
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = _clean("\n".join(parts))
    if not text:
        return "", "empty", "文档里没有可提取的文字（可能整篇都是图片）"
    return text, "ok", ""


def _plain(data: bytes) -> Tuple[str, str, str]:
    for enc in ("utf-8", "gb18030", "utf-16", "latin-1"):
        try:
            text = _clean(data.decode(enc))
            if text:
                return text, "ok", ""
            return "", "empty", "文件是空的"
        except UnicodeDecodeError:
            continue
    return "", "failed", "无法确定文本编码"
